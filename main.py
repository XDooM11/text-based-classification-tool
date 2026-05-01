import pathlib
import pandas as pd
import re
from fpdf import FPDF
from unidecode import unidecode
import pdfplumber
from docx import Document
import os

class FakultniAuditorFinal:
    def __init__(self, target_dir):
        self.target_dir = pathlib.Path(target_dir)
        self.results = []
        self.unmatched = []
        self.all_files_data = [] 
        self.seznam_exts = {'.pdf', '.docx', '.xlsx', '.xls', '.doc'}
        
        # VRÁCENO: Kompletní databáze klíčů včetně všech adres kolejí
        self.fakulty_data = {
            "Cyrilometodejska teologicka fakulta": r"Cyrilometodejsk[at]\s*teologick[at]\s*fakulta|Univerzitni\s*22|\bcmtf\b",
            "Lekarska fakulta": r"Lekarsk[at]\s*fakulta|Hnevotinska\s*3|\blf\b",
            "Filozoficka fakulta": r"Filozofick[at]\s*fakulta|Krizkovskeho\s*10|\bff\b",
            "Prirodovedecka fakulta": r"Prirodovedeck[at]\s*fakulta|17\.\s*listopadu|Slechtitelu|\bprf\b|Prirodoveda",
            "Pedagogicka fakulta": r"Pedagogick[at]\s*fakulta|Zizkovo\s*nam|\bpdf\b",
            "Pravnicka fakulta": r"Pravnick[at]\s*fakulta|17\.\s*listopadu\s*8|\bpf\b",
            "Fakulta telesne kultury": r"Fakulta\s*telesne\s*kultury|tr\.\s*Miru\s*117|\bftk\b",
            "Fakulta zdravotnickych ved": r"Fakulta\s*zdravotnickych\s*ved|Solni\s*19|\bfzv\b",
            "Koleje a menzy (SKM)": r"Katerinska\s*17|Smeralova|17\.\s*listopadu\s*7|Slechtitelu\s*100|J\.\s*L\.\s*Fischera|Bedricha\s*Vaclavka|Evzena\s*Rosickeho|tr\.\s*Miru\s*113|tr\.\s*Miru\s*644|Kuderik[at]|\bskm\b|Koleje\s*a\s*menzy"
        }

    def _get_content(self, file_path):
        ext = file_path.suffix.lower()
        try:
            if ext == ".pdf":
                with pdfplumber.open(file_path) as pdf:
                    return unidecode(" ".join([p.extract_text() or "" for p in pdf.pages[:15]]))
            elif ext == ".docx":
                # VRÁCENO: Načítání Word dokumentů
                doc = Document(file_path)
                return unidecode(" ".join([p.text for p in doc.paragraphs]))
        except:
            return ""
        return ""

    def run_audit(self):
        print(f"--- ANALYZA: {self.target_dir.name} ---")
        for file in self.target_dir.rglob('*'):
            if not file.is_file(): continue
            ext = file.suffix.lower() if file.suffix else "bez pripony"
            size_mb = os.path.getsize(file) / (1024 * 1024)
            
            try:
                rel = file.relative_to(self.target_dir)
                root_folder = rel.parts[0] if len(rel.parts) > 1 else "Hlavni"
            except: root_folder = "Neznamo"

            self.all_files_data.append({'ext': ext, 'size': size_mb, 'folder': root_folder, 'name': file.name})

            if ext in self.seznam_exts:
                obsah = self._get_content(file)
                stem = unidecode(file.stem.lower())
                found_facs = []
                found_trigs = []
                
                for jmeno, reg in self.fakulty_data.items():
                    pattern = unidecode(reg)
                    m_obs = re.search(pattern, obsah, re.IGNORECASE)
                    m_st = re.search(pattern, stem, re.IGNORECASE)
                    
                    if m_obs or m_st:
                        found_facs.append(jmeno)
                        # Uložení klíče, který způsobil shodu
                        t = m_obs.group(0) if m_obs else m_st.group(0)
                        found_trigs.append(t.strip())
                
                if found_facs:
                    self.results.append({
                        "Soubor": unidecode(file.name), 
                        "Fakulty": ", ".join(set(found_facs)),
                        "Klic": ", ".join(set(found_trigs))
                    })
                else:
                    self.unmatched.append(unidecode(file.name))

    def export_pdf(self, output_name="Audit_Manager_Report_Final1.pdf"):
        df = pd.DataFrame(self.all_files_data)
        pdf = FPDF()
        
        # STRANA 1: SOUHRN (Dvě karty + objem dat)
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 18)
        pdf.cell(0, 15, "CELKOVY PREHLED DATASETU", align='C', new_x="LMARGIN", new_y="NEXT")
        pdf.ln(5)

        pdf.set_font("Helvetica", "B", 12)
        pdf.set_fill_color(245, 245, 245)
        pdf.cell(90, 15, f"Celkovy pocet souboru: {len(df)}", border=1, fill=True, align='C')
        pdf.cell(10, 15, "") 
        pdf.cell(90, 15, f"Z toho k auditu: {len(self.results) + len(self.unmatched)}", border=1, fill=True, align='C', new_x="LMARGIN", new_y="NEXT")
        pdf.ln(10)

        pdf.set_font("Helvetica", "I", 10)
        pdf.cell(0, 8, f"Celkovy objem dat: {df['size'].sum():.2f} MB", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(5)

        # STRANA 2: ZAŘAZENÉ (Tabulka s klíčem a fixovaným zarovnáním)
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, f"SEZNAM ZARAZENYCH ({len(self.results)})", align='C', new_x="LMARGIN", new_y="NEXT")
        pdf.ln(5)
        
        w_file, w_fac, w_key = 90, 55, 45
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_fill_color(230, 235, 245)
        pdf.cell(w_file, 8, "Nazev souboru", 1, 0, 'C', True)
        pdf.cell(w_fac, 8, "Fakulta / Soucast", 1, 0, 'C', True)
        pdf.cell(w_key, 8, "Nalezeny klic", 1, 1, 'C', True)
        
        pdf.set_font("Helvetica", "", 7)
        for res in self.results:
            if pdf.get_y() > 260: pdf.add_page()
            s_x, s_y = pdf.get_x(), pdf.get_y()
            pdf.multi_cell(w_file, 5, res["Soubor"], border=1)
            y_f = pdf.get_y()
            pdf.set_xy(s_x + w_file, s_y)
            pdf.multi_cell(w_fac, 5, res["Fakulty"], border=1)
            y_fa = pdf.get_y()
            pdf.set_xy(s_x + w_file + w_fac, s_y)
            pdf.multi_cell(w_key, 5, res["Klic"], border=1)
            y_k = pdf.get_y()
            pdf.set_y(max(y_f, y_fa, y_k))

        # STRANA 3+: NEZAŘAZENÉ
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, f"SEZNAM NEZARAZENYCH ({len(self.unmatched)})", align='C', new_x="LMARGIN", new_y="NEXT")
        pdf.ln(5)
        pdf.set_font("Helvetica", "", 7)
        for name in self.unmatched:
            if pdf.get_y() > 280: pdf.add_page()
            pdf.multi_cell(0, 5, f"- {name}", border=0, new_x="LMARGIN", new_y="NEXT")

        pdf.output(output_name)
        print(f"Hotovo! Report: {os.path.abspath(output_name)}")

if __name__ == "__main__":
    cesta = input("Zadejte cestu ke slozce: ")
    auditor = FakultniAuditorFinal(cesta)
    auditor.run_audit()
    auditor.export_pdf()